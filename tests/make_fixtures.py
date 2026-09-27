"""Regenerate the CV test files in tests/fixtures (pip install python-docx reportlab pikepdf msoffcrypto-tool)."""
import io
import os

import msoffcrypto
import pikepdf
from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

os.makedirs(os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures"), exist_ok=True)
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures"))

CV = """Jordan Lee
jordan.lee@example.com | Manchester

SUMMARY
Customer service professional with 4 years in retail and call-centre roles. Comfortable with spreadsheets and learning new software.

EXPERIENCE
Customer Support Advisor - BrightTel (2022-2025)
• Answered 60+ customer calls per day about billing and technical issues
• Logged tickets in Zendesk and escalated complex cases to tier 2
• Trained 5 new starters on the call scripts and ticketing system

Sales Assistant - HomeGoods Store (2020-2022)
• Handled till, returns and stock counts
• Built a simple Excel sheet to track weekly stock shortages

EDUCATION
BTEC Level 3 Business, City College (2020)

SKILLS
Zendesk, Excel, Microsoft Office, conflict resolution"""


def text_pdf(path):
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica", 11)
    y = 800
    for line in CV.split("\n"):
        c.drawString(50, y, line.replace("•", "-"))
        y -= 16
    c.save()


# 1. Normal text PDF
text_pdf("cv.pdf")

# 2. Normal DOCX (with real bullet style)
doc = Document()
for line in CV.split("\n"):
    if line.startswith("• "):
        doc.add_paragraph(line[2:], style="List Bullet")
    else:
        doc.add_paragraph(line)
doc.save("cv.docx")

# 3. "Scanned" PDF: the CV rendered as a picture, no selectable text
img = Image.new("RGB", (1240, 1754), "white")
d = ImageDraw.Draw(img)
for i, line in enumerate(CV.split("\n")):
    d.text((80, 80 + i * 30), line.replace("•", "-"), fill="black")
buf = io.BytesIO()
img.save(buf, format="PNG")
buf.seek(0)
c = canvas.Canvas("scanned.pdf", pagesize=A4)
c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1])
c.save()

# 4. Password-protected PDF
with pikepdf.open("cv.pdf") as pdf:
    pdf.save("locked.pdf", encryption=pikepdf.Encryption(user="secret", owner="secret"))

# 5. Password-protected DOCX
with open("cv.docx", "rb") as src, open("locked.docx", "wb") as out:
    f = msoffcrypto.OfficeFile(src)
    f.encrypt("secret", out)

# 6. Old .doc (OLE header is all our check looks at)
with open("old.doc", "wb") as f:
    f.write(bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1]) + b"\0" * 5000)

# 7. Oversized files are generated in memory by tests/upload.mjs, not stored in the repo.

# 8. Empty file
open("empty.pdf", "wb").close()

# 9. Not a CV file at all, renamed to .pdf
with open("notes.pdf", "w") as f:
    f.write("just some plain text pretending to be a pdf")

print("created:", sorted(p for p in os.listdir(".") if "." in p and not p.endswith(".py")))
